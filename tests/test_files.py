"""文件解析单元测试。"""
from __future__ import annotations

from src.config import FilesConfig
from src.files import FileSummarizer


def make(cfg_kwargs=None):
    cfg = FilesConfig(**(cfg_kwargs or {}))
    return FileSummarizer(cfg)


def test_text_file(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("今天开会：预算 3000，下周三交。", encoding="utf-8")
    out = make().summarize(p)
    assert "note.txt" in out
    assert "预算 3000" in out


def test_gbk_text_file(tmp_path):
    p = tmp_path / "old.txt"
    p.write_bytes("中文 GBK 内容".encode("gbk"))
    out = make().summarize(p)
    assert "中文 GBK 内容" in out


def test_docx_file(tmp_path):
    from docx import Document

    p = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("第一段：季度总结")
    doc.add_paragraph("第二段：利润增长")
    doc.save(p)
    out = make().summarize(p)
    assert "report.docx" in out
    assert "季度总结" in out


def test_xlsx_file(tmp_path):
    from openpyxl import Workbook

    p = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "姓名"
    ws["B1"] = "分数"
    ws["A2"] = "小明"
    ws["B2"] = 95
    wb.save(p)
    out = make().summarize(p)
    assert "data.xlsx" in out
    assert "小明" in out


def test_binary_file(tmp_path):
    p = tmp_path / "blob.bin"
    p.write_bytes(b"\x00\x01\x02" * 100)
    out = make().summarize(p)
    assert "blob.bin" in out
    assert "二进制" in out


def test_too_large_file(tmp_path):
    p = tmp_path / "big.txt"
    p.write_bytes(b"x" * (2 * 1024 * 1024))
    out = make({"max_file_mb": 1}).summarize(p)
    assert "太大" in out


def test_long_text_truncated(tmp_path):
    p = tmp_path / "long.txt"
    p.write_text("啊" * 5000, encoding="utf-8")
    out = make({"max_text_chars": 100}).summarize(p)
    assert "已截断" in out
