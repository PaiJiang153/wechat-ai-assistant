"""文件消息处理：把群里的文件解析成一段能进上下文的中文描述。

策略
====
* 文本类文件（txt/md/csv/json/code…）：读前 max_text_chars 个字符；
* Word（.docx）：python-docx 提取正文；
* Excel（.xlsx）：openpyxl 读前几行；
* PDF：装了 pypdf 就提取文本，没装就只报文件名和大小；
* 其它二进制/超大文件：只报文件名、类型、大小，不硬读。
"""
from __future__ import annotations

import logging
from pathlib import Path

from .config import FilesConfig

log = logging.getLogger("dsbot.files")

_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".log", ".ini",
    ".cfg", ".conf", ".yaml", ".yml", ".xml", ".html", ".htm", ".sql", ".sh", ".bat",
    ".ps1", ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".kt", ".c", ".h", ".cpp",
    ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".lua", ".m", ".swift",
}


class FileSummarizer:
    def __init__(self, cfg: FilesConfig) -> None:
        self.cfg = cfg

    # ------------------------------------------------------------------
    def summarize(self, path: str | Path, display_name: str = "") -> str:
        """返回形如「报告.docx（12.3 KB）：……节选……」的一句话描述。"""
        p = Path(path)
        name = display_name or p.name
        if not p.exists():
            return ""
        size = p.stat().st_size
        size_text = self._fmt_size(size)
        if not self.cfg.enabled:
            return f"{name}（{size_text}）"
        if size > self.cfg.max_file_mb * 1024 * 1024:
            return f"{name}（{size_text}，文件太大没读内容）"

        ext = p.suffix.lower()
        try:
            if ext in _TEXT_EXTS:
                return f"{name}（{size_text}）：{self._read_text(p)}"
            if ext == ".docx":
                return f"{name}（{size_text}）：{self._read_docx(p)}"
            if ext in (".xlsx", ".xlsm"):
                return f"{name}（{size_text}）：{self._read_xlsx(p)}"
            if ext == ".pdf":
                text = self._read_pdf(p)
                return f"{name}（{size_text}）：{text}" if text else f"{name}（{size_text}，PDF 未装解析库或没有文字层）"
        except Exception as exc:  # noqa: BLE001
            log.warning("解析文件失败 %s：%s", name, exc)
        return f"{name}（{size_text}，二进制文件，没读内容）"

    # ------------------------------------------------------------------
    @staticmethod
    def _fmt_size(size: int) -> str:
        if size < 1024:
            return f"{size} B"
        if size < 1024 * 1024:
            return f"{size / 1024:.1f} KB"
        return f"{size / 1024 / 1024:.1f} MB"

    def _clip(self, text: str) -> str:
        text = " ".join(str(text).split())
        if len(text) > self.cfg.max_text_chars:
            text = text[: self.cfg.max_text_chars] + "……（已截断）"
        return text

    def _read_text(self, p: Path) -> str:
        with p.open("rb") as stream:
            raw = stream.read(max(4096, self.cfg.max_text_chars * 4 + 4))
        for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
            try:
                import codecs
                return self._clip(codecs.getincrementaldecoder(enc)().decode(raw, final=False))
            except UnicodeDecodeError:
                continue
        return "（编码识别失败）"

    def _read_docx(self, p: Path) -> str:
        from docx import Document

        doc = Document(str(p))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
            if sum(len(p) for p in parts) >= self.cfg.max_text_chars:
                return self._clip("\n".join(parts))
        for table in doc.tables[:3]:
            for row in table.rows[:10]:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return self._clip("\n".join(parts)) or "（空文档）"

    def _read_xlsx(self, p: Path) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(str(p), read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets[:2]:
            parts.append(f"[工作表]{ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 10:
                    parts.append("……")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    parts.append(" | ".join(cells))
        wb.close()
        return self._clip("\n".join(parts)) or "（空表格）"

    def _read_pdf(self, p: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError:
            return ""
        reader = PdfReader(str(p))
        parts: list[str] = []
        for page in reader.pages[:5]:
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(text)
        return self._clip("\n".join(parts))
